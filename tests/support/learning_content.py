"""Owned synthetic content fixture, not an import or practice business service."""

import hashlib
from io import BytesIO

from docx import Document as Docx
from docx.oxml import parse_xml
from PIL import Image, ImageDraw
from sqlalchemy.orm import Session

from backend.app.models import Document, Question, QuestionAsset
from backend.app.models.practice import PracticeItem, PracticeSession
from backend.app.models.question_origin import QuestionOrigin
from backend.app.models.question_revision import (
    QuestionRevision,
    question_revision_assets,
)
from backend.app.models.subject import Subject


def registration(data, media_type="image/png"):
    return {"media_type": media_type, "size_bytes": len(data),
            "sha256": hashlib.sha256(data).hexdigest(), "migration": {"status": "not_required"}}


def seed_content(engine, profile, *, configure=None):
    root = profile.storage_root
    (root / "content").mkdir(parents=True, exist_ok=True)
    docx = Docx()
    docx.add_paragraph("DOCX 合成内容：文字、图片、公式和表格")
    paragraph = docx.add_paragraph("x 的平方：")
    paragraph._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath>'))
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "时间(s)", "位移(m)"
    table.cell(1, 0).text, table.cell(1, 1).text = "1", "2"
    images = []
    for label in ("STEM 1", "STEM 2", "OPTION A1", "OPTION A2", "OPTION B1", "OPTION B2", "EXPLANATION 1", "EXPLANATION 2", "ORIGINAL x ? 2"):
        image = Image.new("RGB", (240, 80), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((1, 1, 238, 78), outline="steelblue", width=3)
        draw.text((12, 30), label, fill="black")
        stream = BytesIO()
        image.save(stream, format="PNG")
        data = stream.getvalue()
        images.append(data)
        docx.add_paragraph(label).add_run().add_picture(BytesIO(data))
    stream = BytesIO()
    docx.save(stream)
    source_bytes = stream.getvalue()
    (root / "content/source.docx").write_bytes(source_bytes)
    formula = {"type": "formula", "expression": {"kind": "sup", "children": [
        {"kind": "identifier", "text": "x"}, {"kind": "number", "text": "2"}]}}
    table_block = {"type": "table", "headers": ["时间(s)", "位移(m)"], "rows": [["1", "2"]]}
    with Session(engine) as session, session.begin():
        subject = Subject(name="图文合同合成科目")
        other = Subject(name="其他科目")
        session.add_all([subject, other])
        session.flush()
        document = Document(subject_id=subject.id, original_filename="content.docx", file_format="docx",
                            purpose="paper_source", status="Uploaded", storage_path="content/source.docx",
                            file_metadata=registration(source_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), retryable=False)
        question = Question(subject_id=subject.id, knowledge_points=[], status="Draft")
        session.add_all([document, question])
        session.flush()
        assets = []
        for index, data in enumerate(images):
            locator = f"content/image-{index}.png"
            (root / locator).write_bytes(data)
            asset = QuestionAsset(question_id=question.id, asset_type="figure", width=240, height=80,
                                  file_path=locator, file_metadata=registration(data), order_index=None,
                                  source_document_id=document.id, source_anchor=f"word/document.xml#p[{index + 3}]/r/drawing")
            session.add(asset)
            assets.append(asset)
        session.flush()
        def image(index, label):
            return {"type": "image", "asset_id": str(assets[index].id), "alt": label}
        original = {"type": "original", "asset_id": str(assets[8].id), "reason": "复杂对象待确认，不能直接自动评分",
                    "source": {"document_id": str(document.id), "anchor": "word/document.xml#p[11]/r/drawing"}}
        revision = QuestionRevision(question_id=question.id, subject_id=subject.id, revision=1,
            question_type="single_choice", stem_blocks=[{"type": "text", "text": "文字之前"}, image(0, "题干第一图"),
            {"type": "text", "text": "两图之间"}, image(1, "题干第二图"), formula, table_block, original],
            options=[{"id": "stable-a", "blocks": [{"type": "text", "text": "选项甲"}, image(2, "选项甲第一图"), formula, image(3, "选项甲第二图"), table_block]},
                     {"id": "stable-b", "blocks": [{"type": "text", "text": "选项乙"}, image(4, "选项乙第一图"), table_block, image(5, "选项乙第二图"), formula]}],
            explanation_blocks=[{"type": "text", "text": "整组提交后的解析"}, image(6, "解析第一图"), formula,
                                image(7, "解析第二图"), table_block],
            reference_answer="stable-a", scoring_rubric="PRIVATE RUBRIC", grounding="uploaded_material",
            admission_state="needs_correction")
        session.add(revision)
        session.flush()
        question.current_revision_id = revision.id
        session.execute(question_revision_assets.insert(), [{"revision_id": revision.id, "asset_id": asset.id} for asset in assets])
        origin = QuestionOrigin(question_id=question.id, subject_id=subject.id, source_kind="uploaded", document_id=document.id,
                                source_location={"anchor": original["source"]["anchor"]})
        session.add(origin)
        session.add_all([QuestionOrigin(question_id=question.id, subject_id=subject.id, source_kind="uploaded",
            document_id=document.id, source_location={"anchor": asset.source_anchor}) for asset in assets[:8]])
        session.flush()
        if configure is not None:
            configure(session, question, document, assets, revision)
            session.flush()
        practice = PracticeSession(subject_id=subject.id, state="preview", conditions={})
        session.add(practice)
        session.flush()
        item = PracticeItem(session_id=practice.id, subject_id=subject.id, position=1, question_id=question.id,
                            question_revision_id=revision.id, fixed_max_score=10, scoring_input={"reference_answer": "stable-a"}, knowledge_point_ids=[])
        session.add(item)
        session.flush()
        practice.state = "in_progress"
        return {"subject": subject.id, "other_subject": other.id, "document": document.id, "question": question.id,
                "revision": revision.id, "practice": practice.id, "item": item.id, "origin": origin.id,
                "assets": [a.id for a in assets], "formula": formula, "table": table_block}
