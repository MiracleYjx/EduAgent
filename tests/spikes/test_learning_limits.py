"""T005 resource observations and real Adapter/Fake HTTP transport, NOT product limits."""
from __future__ import annotations
import asyncio
import base64
from io import BytesIO
import json
import os
from pathlib import Path
import threading
import time
import tracemalloc
from types import SimpleNamespace
from uuid import UUID
from zipfile import ZipFile
from docx import Document
from PIL import Image, ImageDraw
from pydantic import BaseModel, Field, SecretStr
from pypdf import PdfWriter
import httpx
import pytest
from openai import AsyncOpenAI
from backend.app.ai.ingestion.parsers import parse_document
from backend.app.ai.ingestion.paper_pipeline import open_paper, PaperInputError
from backend.app.ai.ingestion.ocr.rapidocr import RapidOCRProvider
from backend.app.ai.llm.openai_compatible import OpenAICompatibleProvider
from backend.app.ai.llm.deepseek import DeepSeekProvider
from backend.app.ai.embedding.providers.openai_compatible import OpenAICompatibleEmbeddingProvider
from backend.app.ai.retrieval.base import RetrievedChunk, RetrievalMode
from backend.app.ai.retrieval.reranker import LLMRerankAdapter
from backend.app.ai.agents import question_agent as generation
from backend.app.ai.agents.state import QuestionGenerationRequest
from backend.app.services.grading import grading_context as grading
from backend.app.ai.vision.provider import ProviderVisionUnderstanding
from backend.app.ai.vision.base import VisionFailure, VisionImage
from backend.app.core.retry_policy import RetryPolicy, ProviderExecutionError

OUT=Path(os.environ['PL002_OUTPUT_ROOT'])

def save_result(name,value):
    directory=OUT/'limits'; directory.mkdir(exist_ok=True)
    (directory/(name+'.json')).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
    print(name+': '+json.dumps(value,ensure_ascii=False))


def test_text_bytes_pages_and_image_object_sweep(tmp_path):
    measurements=[]
    for mib in [1,5,20]:
        data=b'synthetic teaching content.\n'*(mib*1024*1024//28+1); data=data[:mib*1024*1024]
        path=tmp_path/f'text-{mib}MiB.txt'; path.write_bytes(data)
        tracemalloc.start(); started=time.perf_counter(); parsed=parse_document(path.name,data)
        _current,peak=tracemalloc.get_traced_memory(); tracemalloc.stop()
        assert parsed.text and parsed.sections
        measurements.append({'kind':'txt','file_bytes':len(data),'seconds':round(time.perf_counter()-started,4),'python_tracemalloc_peak_bytes':peak,'not_rss_or_global_upload_limit':True})
    for pages in [1,10,50,51]:
        writer=PdfWriter()
        for _ in range(pages): writer.add_blank_page(width=595,height=842)
        output=BytesIO(); writer.write(output); data=output.getvalue()
        if pages==51:
            with pytest.raises(PaperInputError) as exc:
                with open_paper(data): pass
            assert exc.value.code=='PAPER_TOO_MANY_PAGES'
            measurements.append({'kind':'pdf','pages':pages,'bytes':len(data),'current_legacy_rejection':exc.value.code,'new_limit_proposed':False})
        else:
            started=time.perf_counter(); decoded=0; dims=None
            with open_paper(data) as paper:
                for page in paper.pages():
                    dims=[page.width,page.height]; decoded+=page.width*page.height*3
            measurements.append({'kind':'pdf','pages':pages,'bytes':len(data),'seconds':round(time.perf_counter()-started,4),'rendered_dimensions':dims,'total_rgb_bytes_if_all_retained':decoded,'streaming_measured_not_ocr':True})
    for count in [1,5,8]:
        doc=Document(); paragraph=doc.add_paragraph('synthetic image-order sweep ')
        for i in range(count):
            image=Image.new('RGB',(128,96),'white'); ImageDraw.Draw(image).text((5,5),str(i),fill='black')
            content=BytesIO(); image.save(content,format='PNG'); content.seek(0); paragraph.add_run().add_picture(content)
        output=BytesIO(); doc.save(output); data=output.getvalue(); started=time.perf_counter(); loaded=Document(BytesIO(data))
        assert len(loaded.inline_shapes)==count
        with ZipFile(BytesIO(data)) as z: media=len([n for n in z.namelist() if n.startswith('word/media/')])
        measurements.append({'kind':'docx','images':count,'media_objects':media,'file_bytes':len(data),'load_seconds':round(time.perf_counter()-started,4),'not_product_image_limit':True})
    save_result('local-sweep',measurements)


class TextResult(BaseModel):
    result: str


class BoundedTextResult(BaseModel):
    result: str = Field(max_length=16)


def completion(content):
    return {'id':'synthetic','object':'chat.completion','created':0,'model':'synthetic-not-a-real-model','choices':[{'index':0,'message':{'role':'assistant','content':content},'finish_reason':'stop'}],'usage':{'prompt_tokens':10,'completion_tokens':5,'total_tokens':15}}


def test_fake_http_real_adapter_captures_output_flag_without_local_hard_token_cap():
    async def run():
        calls=[]
        def handler(request):
            calls.append({'body':json.loads(request.content),'bytes':len(request.content),'timeout':dict(request.extensions.get('timeout',{})),'authorization_present':'authorization' in request.headers})
            return httpx.Response(200,json=completion(json.dumps({'result':'x'*10000})))
        http=httpx.AsyncClient(transport=httpx.MockTransport(handler),trust_env=False)
        provider=OpenAICompatibleProvider(model='synthetic',base_url='https://mock.invalid/v1',auth_mode='none',http_client=http,max_tokens=64,timeout=0.25,retry_policy=RetryPolicy(max_retries=0))
        result=await provider.generate_structured([{'role':'user','content':'JSON: '+('合成上下文'*1000)}],TextResult)
        assert len(result.result)==10000 and len(calls)==1
        assert calls[0]['body']['max_tokens']==64 and calls[0]['body']['response_format']=={'type':'json_object'}
        assert not calls[0]['authorization_present']
        with pytest.raises(ProviderExecutionError):
            await provider.generate_structured([{'role':'user','content':'JSON'}],BoundedTextResult)
        assert len(calls)==2
        await provider.aclose(); assert not http.is_closed; await http.aclose()
        save_result('llm-wire',{'calls':calls,'returned_characters':len(result.result),'fake_usage_not_billing':True,'max_tokens_is_forwarded_not_locally_proven_hard_cap':True})
    asyncio.run(run())


@pytest.mark.parametrize('failure,old_attempts',[('rate_limit',3),('timeout',3),('invalid_json',2)])
def test_fake_http_default_retry_amplification_and_explicit_zero_retry(failure,old_attempts):
    async def run():
        records=[]
        for retries,expected in [(2,old_attempts),(0,1)]:
            calls=[]
            def handler(request):
                calls.append(len(request.content))
                if failure=='timeout': raise httpx.ReadTimeout('synthetic timeout after send',request=request)
                if failure=='rate_limit': return httpx.Response(429,json={'error':{'message':'synthetic','type':'rate_limit_error'}},headers={'retry-after':'0'})
                return httpx.Response(200,json=completion('not JSON'))
            http=httpx.AsyncClient(transport=httpx.MockTransport(handler),trust_env=False)
            policy=RetryPolicy(max_retries=retries,backoff_seconds=(0.0,),jitter_seconds=0.0)
            provider=OpenAICompatibleProvider(model='synthetic',base_url='https://mock.invalid/v1',auth_mode='none',http_client=http,retry_policy=policy)
            with pytest.raises(ProviderExecutionError) as exc: await provider.generate_structured([{'role':'user','content':'JSON'}],TextResult)
            assert len(calls)==expected and exc.value.attempt_count==expected
            records.append({'retry_count':retries,'actual_http_sends':len(calls),'error':exc.value.code,'unknown_billing_on_timeout':failure=='timeout'})
            await provider.aclose(); await http.aclose()
        assert RetryPolicy().max_retries==2
        save_result('retry-'+failure,records)
    asyncio.run(run())


def test_real_embedding_batching_and_rerank_count_actual_http_steps():
    async def run():
        requests=[]
        def handler(request):
            body=json.loads(request.content); requests.append({'path':request.url.path,'body':body,'bytes':len(request.content)})
            if request.url.path.endswith('/embeddings'):
                return httpx.Response(200,json={'object':'list','model':'synthetic','data':[{'object':'embedding','index':i,'embedding':[0.1,0.2,0.3]} for i in range(len(body['input']))],'usage':{'prompt_tokens':0,'total_tokens':0}})
            # Only the first two candidates are sent by the actual reranker.
            return httpx.Response(200,json=completion(json.dumps({'rankings':[{'chunk_id':'chunk-0','score':0.9},{'chunk_id':'chunk-1','score':0.5}]})))
        http=httpx.AsyncClient(transport=httpx.MockTransport(handler),trust_env=False)
        embed=OpenAICompatibleEmbeddingProvider(model='synthetic',api_key='',auth_mode='none',base_url='https://mock.invalid/v1',dimension=3,batch_size=2,max_retries=0,http_client=http)
        vectors=await embed.embed_documents(['synthetic '+str(i) for i in range(5)]); query=await embed.embed_query('query')
        assert len(vectors)==5 and len(query)==3 and len(requests)==4
        provider=OpenAICompatibleProvider(model='synthetic',base_url='https://mock.invalid/v1',auth_mode='none',http_client=http,retry_policy=RetryPolicy(max_retries=0))
        rerank=LLMRerankAdapter(provider=provider,max_candidates=2,timeout=0.5)
        chunks=[RetrievedChunk('chunk-'+str(i),'course','document','synthetic '+('data '*2000)) for i in range(5)]
        ranked=await rerank.rerank_async('query',chunks,top_k=2)
        assert len(requests)==5 and len(ranked)==2
        sent=requests[-1]['body']['messages'][-1]['content']; assert 'chunk-2' not in sent
        await rerank.aclose(); await embed.aclose(); await provider.aclose(); assert not http.is_closed; await http.aclose()
        save_result('derived-http',{'requests':requests,'embedding_documents':5,'experimental_batch_size':2,'embedding_document_calls':3,'query_embedding_calls':1,'rerank_calls':1,'database_vector_identity_not_tested':True})
    asyncio.run(run())


def test_real_context_assemblers_bound_retrieval_but_not_complete_prompt():
    chunks=[RetrievedChunk(str(UUID(int=i+1)),str(UUID(int=100)),str(UUID(int=200)),'教学依据'*2000) for i in range(20)]
    result={}
    for name,module in [('generation',generation),('grading',grading)]:
        written,text=module._compose_final_context(chunks)
        assert len(text)<=module.MAX_FINAL_CONTEXT_CHARS and len(written)<len(chunks)
        assert all(c.chunk_id in text for c in written)
        assert all(c.chunk_id not in text for c in chunks[len(written):])
        result[name]={'retrieval_context_chars':len(text),'utf8_bytes':len(text.encode('utf-8')),'written_chunks':len(written),'current_retrieval_char_budget':module.MAX_FINAL_CONTEXT_CHARS,'per_chunk_char_budget':module.PER_CHUNK_CHARS}
    request=QuestionGenerationRequest(course_id=str(UUID(int=100)),knowledge_points=['长条件'*2000],count=1)
    written,text=generation._compose_final_context(chunks)
    context=generation.QuestionGenerationContext(request,'synthetic query',RetrievalMode.KEYWORD_ONLY,written,tuple(c.chunk_id for c in written),text,len(chunks))
    messages=generation.build_generation_messages(request,context)
    total_chars=sum(len(m['content']) for m in messages)
    assert total_chars>generation.MAX_FINAL_CONTEXT_CHARS
    result['full_generation_messages']={'chars':total_chars,'utf8_bytes':len(json.dumps(messages,ensure_ascii=False).encode('utf-8')),'not_token_count':True,'teacher_context_builder':'not implemented','global_input_token_limit':'not implemented'}
    save_result('context',result)


def test_real_vision_transport_has_old_count_bound_and_fake_http_only():
    async def run():
        calls=[]
        def handler(request):
            body=json.loads(request.content); calls.append({'bytes':len(request.content),'body':body})
            return httpx.Response(200,json=completion(json.dumps({'observations':[],'conditions':[],'unresolved_issues':[],'requires_manual_review':True,'provenance':None})))
        http=httpx.AsyncClient(transport=httpx.MockTransport(handler),trust_env=False)
        # The model/endpoint labels are required solely by the existing Adapter's capability gate.
        # All bytes terminate at MockTransport, NOT a real model or its endpoint.
        sdk=AsyncOpenAI(api_key='synthetic-test-only',base_url='https://api.deepseek.com/v1',http_client=http,max_retries=0)
        settings=SimpleNamespace(deepseek_model='deepseek-v4-flash',deepseek_api_key=SecretStr('synthetic-test-only'),deepseek_base_url='https://api.deepseek.com/v1')
        provider=DeepSeekProvider(settings,client=sdk,retry_policy=RetryPolicy(max_retries=0))
        vision=ProviderVisionUnderstanding(provider); buffer=BytesIO(); Image.new('RGB',(32,32),'white').save(buffer,format='PNG'); data=buffer.getvalue(); image=VisionImage.from_bytes(data)
        assert provider.supports_vision()
        await vision.understand(context={'synthetic':True},images=[image]*5,task='synthetic transport only')
        assert len(calls)==1
        parts=calls[0]['body']['messages'][1]['content']; transported=[part['image_url']['url'] for part in parts if part['type']=='image_url']
        assert len(transported)==5 and base64.b64decode(transported[0].split(',',1)[1])==data
        with pytest.raises(VisionFailure) as exc: await vision.understand(context={},images=[image]*6,task='synthetic')
        assert exc.value.code=='VISION_IMAGE_TRANSPORT_UNAVAILABLE' and len(calls)==1
        save_result('vision-wire',{'actual_http_sends':1,'images_transported':5,'sixth_image_rejected_before_send':True,'request_bytes':calls[0]['bytes'],'raw_five_images_bytes':5*len(data),'cap_is_legacy_not_new_limit':True,'no_real_vision_quality':True})
        await provider.aclose(); await sdk.close(); assert http.is_closed
    asyncio.run(run())


def test_ocr_timeout_does_not_kill_native_worker_and_lock_serializes_calls(tmp_path):
    class EmptyOutput:
        boxes=None; txts=None; scores=None
    active=0; maximum=0; started=threading.Event(); finished=threading.Event(); mutex=threading.Lock()
    def native_engine(image,**kwargs):
        nonlocal active,maximum
        with mutex: active+=1; maximum=max(maximum,active)
        started.set(); time.sleep(0.08)
        with mutex: active-=1
        finished.set(); return EmptyOutput()
    path=tmp_path/'blank.png'; Image.new('RGB',(80,80),'white').save(path)
    provider=RapidOCRProvider(tmp_path); provider._engine=native_engine; provider._output_type=EmptyOutput
    async def run():
        task=asyncio.create_task(provider.extract_text(path))
        while not started.is_set(): await asyncio.sleep(0.001)
        with pytest.raises(TimeoutError): await asyncio.wait_for(task,timeout=0.01)
        assert not finished.is_set()
        before=time.perf_counter(); await provider.extract_text(path); elapsed=time.perf_counter()-before
        assert maximum==1 and finished.is_set()
        save_result('ocr-cancel',{'fake_native_worker_seconds':0.08,'caller_timeout_seconds':0.01,'worker_still_running_after_timeout':True,'second_call_seconds':round(elapsed,4),'max_concurrent_native_workers':maximum,'real_onnx_quality_not_tested_here':True})
    asyncio.run(run())
